#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
项目数据契约校验

用途：在跑 DesignMetrics 之前，先把「输入数据是否满足计算前置条件」查清楚。

为什么需要它：底表构建用的是 left join + fill_null(0)
（见 data_loader.build_analysis_data），图层名或系数key 写错时不会报错，
只会静默按 0 计入，指标会偏小但看不出来。本脚本把这些静默错变成显式报错。

用法：
    python scripts/validate_project_data.py                     # 校验 data/ 下全部城市
    python scripts/validate_project_data.py data/sh             # 校验指定项目目录
    python scripts/validate_project_data.py data/sh --json out.json

退出码：0 = 无高危问题；1 = 有高危问题；2 = 无可用数据
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

BUILDINGS_DIR_NAMES = ("buildings_csv", "Buildings_csv")

# 三个 xlsx 的必需列（从 data_loader.load_related_data 反推）
REQUIRED_COLUMNS = {
    "楼栋属性.xlsx": ["楼号", "类型", "是否奖励楼栋", "是否精装"],
    "系数key.xlsx": ["系数key", "计容", "物理", "销售", "赠送"],
    "补充信息.xlsx": ["指标名称", "数值"],
}

# 各城市规则读取的补充信息键（从 city_rules.py 反推）
# crash=True 表示缺失会直接抛异常（分母为 0），否则静默按 0 参与计算
CITY_INFO_KEYS = {
    "sh": [("用地面积", True), ("不计容配套面积", False), ("电缆夹层面积", False),
           ("人防地下室面积", False), ("地下机动车数量", False), ("门卫", False)],
    "cs": [("用地面积", True), ("商业面积", False), ("计容配套面积", False),
           ("不计容配套面积", False), ("骑楼面积", False), ("地库面积", False),
           ("人防地下室面积", False), ("地下机动车数量", False), ("地面停车", False),
           ("绿地面积", False), ("配套基底面积", False), ("商业基底面积", False)],
    "hf": [("用地面积", True), ("商业面积", False), ("计容配套面积", False),
           ("不计容配套面积", False), ("地库面积", False), ("人防地下室面积", False),
           ("地下机动车数量", False), ("地面停车", False), ("绿地面积", False),
           ("配套基底面积", False), ("商业基底面积", False)],
    "gy": [("用地面积", True), ("地下停车场面积", False), ("人防地下室面积", False),
           ("地下机动车数量", False)],
}

CITY_ALIAS = {"sh": "sh", "上海": "sh", "cs": "cs", "长沙": "cs",
              "hf": "hf", "合肥": "hf", "gy": "gy", "贵阳": "gy"}


def info_keys_for(city_key: str):
    """城市键回退：sh_change 沿用 sh 的规则；找不到返回空清单"""
    if city_key in CITY_INFO_KEYS:
        return CITY_INFO_KEYS[city_key]
    return CITY_INFO_KEYS.get(city_key.split("_")[0], [])


# 程序读取的系数列；超出这些列的字段目前没有被任何代码使用
KEY_USED_COLUMNS = ["系数key", "计容", "物理", "销售", "赠送"]


def find_buildings_dir(project: Path):
    for name in BUILDINGS_DIR_NAMES:
        candidate = project / name
        if candidate.is_dir():
            return candidate
    return None


def read_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh)), fh


def read_table(path: Path):
    """
    读取 Excel 首个工作表，返回 (列名列表, 行列表)。
    优先用 polars（项目已有依赖，.venv 内即可用），退回 openpyxl。
    两者都不可用时返回 (None, None)。
    """
    try:
        import polars as pl

        df = pl.read_excel(path)
        return [str(c).strip() for c in df.columns], df.rows()
    except Exception:
        pass
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, None
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        wb.close()
        return [], []
    columns = [str(c).strip() if c is not None else "" for c in header]
    rows = list(it)
    wb.close()
    return columns, rows


def column_values(columns, rows, name):
    """取某列的非空值集合；列不存在返回 None"""
    if not columns or name not in columns:
        return None
    idx = columns.index(name)
    return {str(r[idx]).strip() for r in rows if len(r) > idx and r[idx] is not None}


def validate_project(project: Path, city_key: str) -> dict:
    result = {
        "项目目录": str(project),
        "城市": city_key,
        "高危": [],
        "警告": [],
        "提示": [],
        "可校验项": [],
        "统计": {},
    }

    if not project.is_dir():
        result["高危"].append(f"目录不存在：{project}")
        return result

    buildings_dir = find_buildings_dir(project)
    if buildings_dir is None:
        result["高危"].append("缺少 Buildings_csv 目录")
        return result

    csv_files = sorted(buildings_dir.glob("*.csv"))
    if not csv_files:
        result["高危"].append(f"{buildings_dir} 下没有 CSV 文件")
        return result

    # ---------- A. 面积线 CSV ----------
    bad_header, bad_layer, bad_area, empty_csv = [], [], [], []
    layer_keys, layer_btypes, layer_floors = set(), set(), set()
    floor_missing = {}
    no_first_floor = []
    total_rows, total_area = 0, 0.0
    per_building = {}

    for path in csv_files:
        rows, _ = read_csv(path)
        if not rows:
            empty_csv.append(path.name)
            continue
        header = list(rows[0].keys())
        if "图层" not in header or "面积" not in header:
            bad_header.append(f"{path.name}(列：{header})")
            continue
        b_area, b_rows = 0.0, 0
        b_floors = set()
        for row in rows:
            layer = (row.get("图层") or "").strip()
            parts = layer.split("-")
            if len(parts) not in (3, 4):
                bad_layer.append(f"{path.name}:{layer}(段数{len(parts)})")
                continue
            if any(p.strip() == "" for p in parts):
                bad_layer.append(f"{path.name}:{layer}(有空段)")
                continue
            if len(parts) == 4:
                layer_btypes.add(parts[1]); layer_floors.add(parts[2]); layer_keys.add(parts[3])
                b_floors.add(parts[2].upper())
            else:
                layer_btypes.add(parts[1]); layer_floors.add("--"); layer_keys.add(parts[2])
                floor_missing[path.stem] = floor_missing.get(path.stem, 0) + 1
            try:
                area = float(row.get("面积") or 0)
            except (TypeError, ValueError):
                bad_area.append(f"{path.name}:{layer}(面积={row.get('面积')!r})")
                continue
            if area <= 0:
                bad_area.append(f"{path.name}:{layer}(面积={area})")
            b_area += area
            b_rows += 1
            total_rows += 1
        total_area += b_area
        per_building[path.stem] = {"行数": b_rows, "面积": round(b_area, 2)}
        if "1F" not in b_floors:
            no_first_floor.append(path.stem)

    result["统计"] = {
        "楼栋数": len(csv_files),
        "面积行数": total_rows,
        "面积合计": round(total_area, 2),
        "楼型": sorted(layer_btypes),
        "系数key": sorted(layer_keys),
        "楼层样例": sorted(layer_floors)[:15],
        "三层图层楼栋数": len(floor_missing),
    }
    if bad_header:
        result["高危"].append(f"{len(bad_header)} 个 CSV 缺 图层/面积 列：{bad_header[:5]}")
    if bad_layer:
        result["高危"].append(f"{len(bad_layer)} 行图层命名不合法（应为 AA/AB-楼型-[楼层-]系数key）：{bad_layer[:5]}")
    if bad_area:
        result["警告"].append(f"{len(bad_area)} 行面积异常（≤0 或非数值）：{bad_area[:5]}")
    if empty_csv:
        result["警告"].append(f"{len(empty_csv)} 个 CSV 无数据行：{empty_csv[:5]}")
    if floor_missing:
        result["提示"].append(
            f"{len(floor_missing)} 个楼栋含三层图层（楼层记为 --），共 {sum(floor_missing.values())} 行。"
            "这类行不参与 floor=1f 的查询；通常为按楼栋汇总的分项，属出图规范")
    if no_first_floor:
        result["警告"].append(
            f"{len(no_first_floor)} 个楼栋完全没有 1F 图层行，"
            f"不会计入依赖 floor=1f 的指标（如建筑基底面积）：{no_first_floor[:10]}")
    result["可校验项"].append(f"面积线 CSV：{len(csv_files)} 个文件 / {total_rows} 行")

    # ---------- B. 三个 xlsx ----------
    missing_files = []
    columns_map = {}
    for fname, required in REQUIRED_COLUMNS.items():
        fpath = project / fname
        if not fpath.is_file():
            missing_files.append(fname)
            continue
        columns, rows = read_table(fpath)
        if columns is None:
            result["提示"].append(f"{fname} 无法读取（polars 与 openpyxl 均不可用），跳过列校验")
            continue
        columns_map[fname] = columns
        lack = [c for c in required if c not in columns]
        if lack:
            result["高危"].append(f"{fname} 缺必需列：{lack}（现有列：{columns}）")
        else:
            result["提示"].append(f"{fname} 列名齐全，{len(rows)} 行数据")

    if missing_files:
        result["高危"].append(
            f"缺输入文件：{missing_files}（这三个文件被 .gitignore 的 *.xlsx 排除，需从内网单独提供）")
        result["提示"].append("缺 xlsx 时无法做交叉校验，也无法运行 DesignMetrics")
        return result

    # ---------- C. 交叉校验 ----------
    key_path = project / "系数key.xlsx"
    key_columns, key_rows = read_table(key_path)
    excel_keys = column_values(key_columns, key_rows, "系数key")
    if excel_keys is not None:
        unknown = sorted(layer_keys - excel_keys)
        unused = sorted(excel_keys - layer_keys)
        extra_cols = [c for c in key_columns if c not in KEY_USED_COLUMNS]
        if unknown:
            result["提示"].append(
                f"面积线里有 {len(unknown)} 个系数key 未在 系数key.xlsx 中定义，这些行按 0 计：{unknown}。"
                "若它们是「主体」的子项分项（如 套内），按 0 计才正确，否则会重复计算；请逐项确认口径")
        else:
            result["可校验项"].append(f"系数key 全部命中（{len(layer_keys)} 个）")
        if unused:
            result["提示"].append(f"系数key.xlsx 中有 {len(unused)} 个 key 未被任何面积线使用：{unused}")
        if extra_cols:
            result["提示"].append(
                f"系数key.xlsx 含程序未读取的列 {extra_cols}（程序只读 {KEY_USED_COLUMNS}），对应口径尚未实现")

    attr_path = project / "楼栋属性.xlsx"
    attr_columns, attr_rows = read_table(attr_path)
    excel_units = column_values(attr_columns, attr_rows, "楼号")
    if excel_units is not None:
        csv_units = set(per_building)
        orphan_csv = sorted(csv_units - excel_units)
        orphan_attr = sorted(excel_units - csv_units)
        if orphan_csv:
            result["高危"].append(
                f"{len(orphan_csv)} 个楼栋在 楼栋属性.xlsx 中无记录，"
                f"类型/是否奖励楼栋/精装会为空：{orphan_csv[:10]}")
        else:
            result["可校验项"].append(f"楼号全部命中（{len(csv_units)} 个楼栋）")
        if orphan_attr:
            result["警告"].append(
                f"楼栋属性.xlsx 中有 {len(orphan_attr)} 个楼号没有对应面积线，"
                f"其户配仍全额计入户数，会导致「总户数」与「建筑面积」口径不一致：{orphan_attr[:10]}")

    # ---------- D. 城市规则必需键 ----------
    info_path = project / "补充信息.xlsx"
    info_columns, info_rows = read_table(info_path)
    info_keys = column_values(info_columns, info_rows, "指标名称")
    if info_keys is not None:
        required = info_keys_for(city_key)
        crash_missing = [k for k, crash in required if crash and k not in info_keys]
        zero_missing = [k for k, crash in required if not crash and k not in info_keys]
        if crash_missing:
            result["高危"].append(
                f"补充信息.xlsx 缺 {crash_missing}，{city_key} 规则会因分母为 0 直接抛异常")
        if zero_missing:
            result["警告"].append(
                f"补充信息.xlsx 缺 {len(zero_missing)} 个 {city_key} 规则读取的键，"
                f"对应指标会静默按 0 计：{zero_missing}")
        if not crash_missing and not zero_missing:
            result["可校验项"].append(
                f"补充信息键齐全（{len(required)} 个）" if required
                else "该目录无对应城市规则键清单，跳过必需键校验")
    else:
        result["提示"].append("补充信息.xlsx 无 指标名称 列或无法读取，跳过必需键校验")

    return result


def main() -> int:
    ap = argparse.ArgumentParser(description="building_metrics 项目数据契约校验")
    ap.add_argument("paths", nargs="*", help="项目目录，缺省校验 data/ 下全部城市")
    ap.add_argument("--json", help="把结果写入 JSON 文件")
    args = ap.parse_args()

    root = Path(__file__).resolve().parents[1]
    if args.paths:
        targets = [(city_key_of(Path(p)), Path(p)) for p in args.paths]
    else:
        data_dir = root / "data"
        if not data_dir.is_dir():
            print(f"未找到数据目录：{data_dir}")
            return 2
        targets = [(p.name, p) for p in sorted(data_dir.iterdir()) if p.is_dir()]

    if not targets:
        print("没有可校验的项目目录")
        return 2

    results = []
    high_total = 0
    runnable = 0
    for city_key, project in targets:
        res = validate_project(project, city_key)
        results.append(res)
        high_total += len(res["高危"])
        if not res["高危"]:
            runnable += 1
        print(f"\n=== {project}（{city_key}） ===")
        stat = res["统计"]
        if stat:
            print(f"  楼栋 {stat['楼栋数']} 个 / 面积行 {stat['面积行数']} / 面积合计 {stat['面积合计']:,}")
            print(f"  楼型：{stat['楼型']}")
            print(f"  系数key（{len(stat['系数key'])}）：{stat['系数key']}")
        for item in res["可校验项"]:
            print(f"  [通过] {item}")
        for item in res["高危"]:
            print(f"  [高危] {item}")
        for item in res["警告"]:
            print(f"  [警告] {item}")
        for item in res["提示"]:
            print(f"  [提示] {item}")

    print(f"\n== 汇总 ==\n可校验项目 {len(results)} 个，其中无高危问题 {runnable} 个；高危问题合计 {high_total} 条")

    if args.json:
        Path(args.json).write_text(
            json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"结果已写入 {args.json}")

    if high_total:
        return 1
    return 0


def city_key_of(path: Path) -> str:
    return CITY_ALIAS.get(path.name, path.name)


if __name__ == "__main__":
    sys.exit(main())
